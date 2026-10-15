"""Exports PDF et Excel des écrans « Rapports » et « Historique » (15/10/2026).

Un rapport (ou l'historique filtré) est décrit une seule fois, sous une forme commune :

    {"titre": "État du stock", "sous_titre": "Magasin : tous · au 15/10/2026",
     "colonnes": [{"cle": "code", "titre": "Code", "type": "code", "largeur": 22}, ...],
     "sections": [{"titre": "Ankorondrano", "lignes": [{...}], "total": {...} ou None}],
     "total": {...} ou None, "resume": [["Articles", "23"], ...], "notes": ["…"]}

Types de colonne : texte, code, quantite (unité dans ligne["_unite"]), nombre, montant (Ar), date (AAAA-MM-JJ),
dateheure (ISO), pill (libellé ; couleur dans ligne["_ton_<cle>"] : rouge, orange, vert, bleu, gris, violet).
`_ton_<cle>` colore aussi le texte des autres colonnes. L'écran affiche la même description (Rapports.jsx).
"""
import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.utils import timezone

LOGO = Path(__file__).resolve().parent / "ressources" / "logo-hbc.png"
MESSAGE_SANS_REPORTLAB = ("Le PDF ne peut pas être fabriqué : le module ReportLab n'est pas installé sur ce PC. "
                          "Double-cliquez sur outils\\installer-modules.bat (Internet nécessaire), puis redémarrez GestiStock.")
MESSAGE_SANS_OPENPYXL = ("Le fichier Excel ne peut pas être fabriqué : le module openpyxl n'est pas installé sur ce PC. "
                         "Double-cliquez sur outils\\installer-modules.bat (Internet nécessaire), puis redémarrez GestiStock.")

TONS = {"rouge": "#b8322a", "orange": "#9a5c00", "vert": "#13785f", "bleu": "#1d5a8c", "gris": "#626262",
        "violet": "#6b46c1"}
FONDS = {"rouge": "#fce4e2", "orange": "#fdf0d8", "vert": "#e3f4ee", "bleu": "#e3f0fb", "gris": "#eef0f2",
         "violet": "#efeafb"}


# ---------------------------------------------------------------- Mise en forme (la même qu'à l'écran)
def nombre(n):
    """6000 → « 6 000 » ; 12.5 → « 12,50 » (espaces entre les milliers, comme à l'écran)."""
    if n is None or n == "":
        return ""
    d = Decimal(str(n))
    if d == d.to_integral_value():
        return f"{int(d):,}".replace(",", " ")
    entier, decimales = f"{d:,.2f}".split(".")
    return f"{entier.replace(',', ' ')},{decimales}"


def quantite(n, unite=None):
    if n is None or n == "":
        return ""
    if unite == "m":
        return f"{nombre(n)} m"
    if unite == "paire":
        return f"{nombre(n)} {'paires' if abs(int(n)) > 1 else 'paire'}"
    return nombre(n)


def montant(v):
    return "" if v is None or v == "" else f"{nombre(Decimal(str(v)).quantize(Decimal('1')))} Ar"


def date_fr(iso):
    if not iso:
        return ""
    if isinstance(iso, (datetime.date, datetime.datetime)):
        iso = iso.isoformat()
    a, m, j = str(iso)[:10].split("-")
    return f"{j}/{m}/{a}"


def dateheure_fr(iso):
    if not iso:
        return ""
    dt = datetime.datetime.fromisoformat(str(iso)) if not isinstance(iso, datetime.datetime) else iso
    if timezone.is_aware(dt):
        dt = timezone.localtime(dt)
    return f"{dt:%d/%m/%Y} {dt:%H:%M}"


def texte_cellule(col, ligne):
    """Valeur affichée d'une cellule (PDF, et Excel pour les colonnes de texte)."""
    v = ligne.get(col["cle"])
    t = col.get("type", "texte")
    if v is None or v == "":
        return ""
    if t == "quantite":
        return quantite(v, ligne.get("_unite"))
    if t == "nombre":
        return nombre(v)
    if t == "montant":
        return montant(v)
    if t == "date":
        return date_fr(v)
    if t == "dateheure":
        return dateheure_fr(v)
    return str(v)


def nom_fichier(titre, extension):
    """« État du stock » → « etat-du-stock-2026-10-15.pdf »."""
    import unicodedata
    base = unicodedata.normalize("NFD", titre.lower())
    base = "".join(c for c in base if not unicodedata.combining(c))
    base = "".join(c if c.isalnum() else "-" for c in base)
    while "--" in base:
        base = base.replace("--", "-")
    return f"{base.strip('-')}-{timezone.localdate().isoformat()}.{extension}"


def _edite(par):
    t = timezone.localtime()
    qui = (par.get_full_name() or par.username) if par else ""
    return f"Édité le {t:%d/%m/%Y} à {t.hour} h {t.minute:02d}" + (f" par {qui}" if qui else "")


# ---------------------------------------------------------------- PDF (ReportLab, A4 paysage)
def fabriquer_pdf(rapport, par=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    bleu, bleu_clair, gris, trait = (colors.HexColor("#2f4763"), colors.HexColor("#e8eef5"), colors.HexColor("#626262"),
                                     colors.HexColor("#b9c2cd"))
    police, gras = "Helvetica", "Helvetica-Bold"
    s_titre = ParagraphStyle("t", fontName=gras, fontSize=15, leading=19, textColor=bleu)
    s_sous = ParagraphStyle("s", fontName=police, fontSize=9.5, leading=12.5, textColor=gris)
    s_section = ParagraphStyle("sec", fontName=gras, fontSize=10.5, leading=13, textColor=bleu, spaceBefore=8, spaceAfter=3)
    s_cel = ParagraphStyle("c", fontName=police, fontSize=8.5, leading=10.5)
    s_cel_d = ParagraphStyle("cd", parent=s_cel, alignment=TA_RIGHT)
    s_ent = ParagraphStyle("e", fontName=gras, fontSize=8.5, leading=10.5, textColor=bleu)
    s_ent_d = ParagraphStyle("ed", parent=s_ent, alignment=TA_RIGHT)
    s_note = ParagraphStyle("n", fontName="Helvetica-Oblique", fontSize=8, leading=10, textColor=gris)
    s_resume = ParagraphStyle("r", fontName=police, fontSize=9.5, leading=12)

    def x(t):
        # Le signe moins typographique (−) n'existe pas dans la police Helvetica du PDF : trait d'union à la place.
        return str(t or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("−", "-")

    a_droite = {"quantite", "nombre", "montant"}
    page = landscape(A4)
    marges = 12 * mm
    largeur = page[0] - 2 * marges
    tampon = BytesIO()
    doc = SimpleDocTemplate(tampon, pagesize=page, leftMargin=marges, rightMargin=marges, topMargin=11 * mm,
                            bottomMargin=14 * mm, title=rapport["titre"], author="HBC – GestiStock")
    el = []
    logo = Image(str(LOGO), width=34 * mm, height=34 * mm * 232 / 746) if LOGO.exists() else Spacer(1, 1)
    tete = Table([[logo, [Paragraph(x(rapport["titre"]), s_titre), Paragraph(x(rapport.get("sous_titre", "")), s_sous)],
                   Paragraph(x(_edite(par)), ParagraphStyle("ed", parent=s_sous, alignment=TA_RIGHT, fontSize=8.5))]],
                 colWidths=[40 * mm, largeur - 100 * mm, 60 * mm])
    tete.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("LINEBELOW", (0, 0), (-1, 0), 1.2, bleu),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    el.append(tete)
    if rapport.get("resume"):
        el.append(Spacer(1, 5))
        el.append(Paragraph("    ".join(f"{x(l)} : <b>{x(v)}</b>" for l, v in rapport["resume"]), s_resume))
    el.append(Spacer(1, 4))

    cols = rapport["colonnes"]
    poids = [c.get("largeur", 20) for c in cols]
    larg = [largeur * p / sum(poids) for p in poids]

    def ligne_pdf(ligne, gras_=False):
        cellules = []
        for c in cols:
            texte = x(texte_cellule(c, ligne))
            ton = ligne.get(f"_ton_{c['cle']}")
            if ton and texte:
                texte = f"<font color='{TONS.get(ton, '#383838')}'><b>{texte}</b></font>"
            if gras_ and texte:
                texte = f"<b>{texte}</b>"
            detail = ligne.get(f"_detail_{c['cle']}")
            if detail:
                texte += f"<br/><font size='7' color='#626262'>{x(detail)}</font>"
            cellules.append(Paragraph(texte, s_cel_d if c.get("type") in a_droite else s_cel))
        return cellules

    def tableau(lignes, total):
        donnees = [[Paragraph(x(c["titre"]), s_ent_d if c.get("type") in a_droite else s_ent) for c in cols]]
        donnees += [ligne_pdf(l) for l in lignes]
        if total:
            donnees.append(ligne_pdf(total, gras_=True))
        t = Table(donnees, colWidths=larg, repeatRows=1)
        style = [("GRID", (0, 0), (-1, -1), 0.5, trait), ("BACKGROUND", (0, 0), (-1, 0), bleu_clair),
                 ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3.5),
                 ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5), ("LEFTPADDING", (0, 0), (-1, -1), 4),
                 ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
        if total:
            style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#f3f5f8")))
        for i, l in enumerate(lignes, start=1):
            if l.get("_fond"):
                style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor(FONDS.get(l["_fond"], "#ffffff"))))
        t.setStyle(TableStyle(style))
        return t

    vide = True
    for s in rapport["sections"]:
        if s.get("titre"):
            el.append(Paragraph(x(s["titre"]), s_section))
        if s["lignes"]:
            vide = False
            el.append(tableau(s["lignes"], s.get("total")))
        else:
            el.append(Paragraph(x(s.get("vide") or "Rien à afficher."), s_note))
    if rapport.get("total") and len(rapport["sections"]) > 1:
        el.append(Spacer(1, 6))
        t = Table([ligne_pdf(rapport["total"], gras_=True)], colWidths=larg)
        t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), 0.8, bleu), ("BACKGROUND", (0, 0), (-1, -1), bleu_clair),
                               ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
        el.append(t)
    if vide and not rapport["sections"]:
        el.append(Paragraph(x(rapport.get("vide") or "Rien à afficher."), s_note))
    for n in rapport.get("notes", []):
        el.append(Spacer(1, 3))
        el.append(Paragraph(x(n), s_note))

    pied = f"GestiStock – {rapport['titre']}"

    def numeroter(canvas, document):
        canvas.saveState()
        canvas.setFont(police, 7.5)
        canvas.setFillColor(gris)
        canvas.drawString(marges, 8 * mm, pied)
        canvas.drawRightString(page[0] - marges, 8 * mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(el, onFirstPage=numeroter, onLaterPages=numeroter)
    return tampon.getvalue()


# ---------------------------------------------------------------- Excel (openpyxl)
def fabriquer_xlsx(rapport, par=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = rapport["titre"][:31].replace("/", "-")
    cols = rapport["colonnes"]
    n = len(cols)
    bleu = "2F4763"
    fin = Side(style="thin", color="B9C2CD")
    bord = Border(left=fin, right=fin, top=fin, bottom=fin)

    ws.cell(row=1, column=1, value=f"HBC – GestiStock : {rapport['titre']}").font = Font(bold=True, size=14, color=bleu)
    ws.cell(row=2, column=1, value=rapport.get("sous_titre", "")).font = Font(italic=True, color="626262")
    ws.cell(row=3, column=1, value=_edite(par)).font = Font(size=9, color="626262")
    r = 4
    if rapport.get("resume"):
        ws.cell(row=r, column=1, value="   ".join(f"{l} : {v}" for l, v in rapport["resume"])).font = Font(bold=True)
        r += 1
    r += 1

    def format_nombre(c, ligne):
        t = c.get("type")
        if t == "montant":
            return '#,##0" Ar"'
        if t == "quantite":
            u = ligne.get("_unite")
            return '#,##0" m"' if u == "m" else '#,##0" paire(s)"' if u == "paire" else "#,##0"
        if t == "nombre":
            return "#,##0"
        return None

    def ecrire(ligne, ligne_n, gras_=False, fond=None):
        for j, c in enumerate(cols, start=1):
            t = c.get("type")
            v = ligne.get(c["cle"])
            if t in ("quantite", "nombre", "montant") and v not in (None, "") and not isinstance(v, str):
                valeur = float(v) if isinstance(v, Decimal) else v
            elif t == "date" and v:
                valeur = datetime.date.fromisoformat(str(v)[:10])
            elif t == "dateheure" and v:
                dt = datetime.datetime.fromisoformat(str(v))
                valeur = (timezone.localtime(dt) if timezone.is_aware(dt) else dt).replace(tzinfo=None)
            else:
                valeur = texte_cellule(c, ligne)
                detail = ligne.get(f"_detail_{c['cle']}")
                if detail:
                    valeur = f"{valeur}\n{detail}" if valeur else detail
            cel = ws.cell(row=ligne_n, column=j, value=valeur if valeur != "" else None)
            cel.border = bord
            cel.alignment = Alignment(vertical="top", wrap_text=True,
                                      horizontal="right" if t in ("quantite", "nombre", "montant") else None)
            fmt = format_nombre(c, ligne)
            if fmt and isinstance(valeur, (int, float)):
                cel.number_format = fmt
            elif t == "date" and v:
                cel.number_format = "DD/MM/YYYY"
            elif t == "dateheure" and v:
                cel.number_format = "DD/MM/YYYY HH:MM"
            ton = ligne.get(f"_ton_{c['cle']}")
            cel.font = Font(bold=gras_ or bool(ton), color=TONS[ton].lstrip("#") if ton in TONS else None)
            if fond:
                cel.fill = PatternFill("solid", fgColor=fond)

    for s in rapport["sections"]:
        if s.get("titre"):
            ws.cell(row=r, column=1, value=s["titre"]).font = Font(bold=True, size=12, color=bleu)
            r += 1
        for j, c in enumerate(cols, start=1):
            cel = ws.cell(row=r, column=j, value=c["titre"])
            cel.font = Font(bold=True, color=bleu)
            cel.fill = PatternFill("solid", fgColor="E8EEF5")
            cel.border = bord
            cel.alignment = Alignment(wrap_text=True, vertical="center")
        r += 1
        if not s["lignes"]:
            ws.cell(row=r, column=1, value=s.get("vide") or "Rien à afficher.").font = Font(italic=True, color="626262")
            r += 1
        for l in s["lignes"]:
            ecrire(l, r, fond=FONDS[l["_fond"]].lstrip("#") if l.get("_fond") in FONDS else None)
            r += 1
        if s.get("total"):
            ecrire(s["total"], r, gras_=True, fond="F3F5F8")
            r += 1
        r += 1
    if rapport.get("total") and len(rapport["sections"]) > 1:
        ecrire(rapport["total"], r, gras_=True, fond="E8EEF5")
        r += 2
    for note in rapport.get("notes", []):
        ws.cell(row=r, column=1, value=note).font = Font(italic=True, size=9, color="626262")
        r += 1

    for j, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(60, c.get("largeur", 20) * 0.9))
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    tampon = BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def reponse_fichier(rapport, format_, par=None):
    """Réponse HTTP : le PDF ou le fichier Excel du rapport (message clair si le module manque)."""
    from django.http import HttpResponse
    from rest_framework import status
    from rest_framework.response import Response
    if format_ == "pdf":
        try:
            contenu = fabriquer_pdf(rapport, par)
        except ImportError:
            return Response({"general": [MESSAGE_SANS_REPORTLAB]}, status=status.HTTP_503_SERVICE_UNAVAILABLE)
        type_, ext = "application/pdf", "pdf"
    elif format_ == "xlsx":
        try:
            contenu = fabriquer_xlsx(rapport, par)
        except ImportError:
            return Response({"general": [MESSAGE_SANS_OPENPYXL]}, status=status.HTTP_503_SERVICE_UNAVAILABLE)
        type_, ext = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        return Response({"general": ["Format inconnu (pdf ou xlsx)."]}, status=status.HTTP_400_BAD_REQUEST)
    reponse = HttpResponse(contenu, content_type=type_)
    reponse["Content-Disposition"] = f'attachment; filename="{nom_fichier(rapport["titre"], ext)}"'
    reponse["Cache-Control"] = "no-store"
    return reponse
