"""Tests des écrans « Générer Rapports » et « Historique » (15/10/2026).
Lancer : python manage.py test stock"""
import datetime
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.core.management import call_command
from django.db.models import Count
from django.test import TestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient

from stock.models import Article, Chantier, Demande, Exemplaire, Magasin, MouvementEtat
from stock.tests_outils import SIGNATURE
from stock.views import articles_avec_quantites
from stock.views_tableau import valeur_du_stock

RAPPORTS = ["etat-stock", "a-commander", "consommation", "materiel-prete", "valeur-stock", "etalonnages"]


@override_settings(GESTISTOCK_EMAILS=False)
class RapportsEtHistoriqueTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("charger_exemples", verbosity=0)
        cls.ank = Magasin.objects.get(nom="Ankorondrano")
        cls.saisisseur = User.objects.create_user("rabe", password="x", first_name="Jean", last_name="Rabe")
        cls.agent = User.objects.create_user("agent", password="x", first_name="Hery", last_name="Rasoa")
        cls.agent.groups.add(Group.objects.get(name="Agent Logistique"))
        cls.resp = User.objects.create_user("resp", password="x", first_name="Paul", last_name="Randria")
        cls.resp.groups.add(Group.objects.get(name="Responsable HBC"))
        cls.admin = User.objects.create_superuser("chef", password="x")
        cls.conso = Article.objects.filter(type="CONSOMMABLE", emplacement__magasin=cls.ank).order_by("code").first()
        Article.objects.filter(id=cls.conso.id).update(qte_stock=50, qte_minimale=10, prix_unitaire=Decimal("2500"))
        cls.casque = Article.objects.create(nom="Casque de chantier", categorie=cls.conso.categorie, type="PRETE",
                                            emplacement=cls.conso.emplacement, qte_stock=10, prix_unitaire=Decimal("30000"))
        cls.serie = (Article.objects.filter(type="PRETE", emplacement__magasin=cls.ank)
                     .annotate(n=Count("exemplaires")).filter(n__gt=0).first())
        for i in range(2):
            Exemplaire.objects.create(article=cls.serie, numero_serie=f"RP-SN-{i}", etat="BON")
        cls.chantier = Chantier.objects.create(nom="Fibre Ivato")

    def setUp(self):
        self.s = APIClient(); self.s.force_login(self.saisisseur)
        self.a = APIClient(); self.a.force_login(self.agent)
        self.r = APIClient(); self.r.force_login(self.resp)

    def circuit(self, retour_prevu=None):
        """Réception, demande (chantier Fibre Ivato), validation, visa, impression, sortie confirmée."""
        rep = self.a.post("/api/entrees/", {"date_reception": timezone.localdate().isoformat(), "magasin": self.ank.id,
                                            "lignes": [{"article": self.conso.id, "quantite": 6, "etat": "NEUF"}]},
                          format="json")
        self.assertEqual(rep.status_code, 201, rep.data)
        retour = (retour_prevu or timezone.localdate() + datetime.timedelta(days=5)).isoformat()
        rep = self.s.post("/api/demandes/", {"magasin": self.ank.id, "chantier": self.chantier.id, "lignes": [
            {"article": self.conso.id, "quantite": 4},
            {"article": self.casque.id, "quantite": 3, "date_retour_prevue": retour},
            {"article": self.serie.id, "quantite": 2, "date_retour_prevue": retour}]}, format="json")
        self.assertEqual(rep.status_code, 201, rep.data)
        d = Demande.objects.get(id=rep.data["id"])
        accord = [{"id": l.id, "qte_accordee": l.quantite} for l in d.lignes.all()]
        self.assertEqual(self.a.post(f"/api/a-valider/{d.id}/valider/", {"signature": SIGNATURE, "mot_de_passe": "x", "lignes": accord},
                                     format="json").status_code, 200)
        self.assertEqual(self.r.post(f"/api/a-viser/{d.id}/viser/", {"signature": SIGNATURE, "mot_de_passe": "x"}, format="json").status_code, 200)
        self.assertEqual(self.a.get(f"/api/requisitions/{d.id}/pdf/").status_code, 200)
        ls = {l.article_id: l for l in d.lignes.all()}
        ids = list(Exemplaire.objects.filter(numero_serie__startswith="RP-SN-").values_list("id", flat=True))
        rep = self.a.post(f"/api/sorties/{d.id}/confirmer/", {"date_remise": timezone.localdate().isoformat(), "lignes": [
            {"id": ls[self.conso.id].id, "qte_livree": 4}, {"id": ls[self.casque.id].id, "qte_livree": 3},
            {"id": ls[self.serie.id].id, "exemplaires": ids}]}, format="json")
        self.assertEqual(rep.status_code, 200, rep.data)
        return d

    # ---------------------------------------------------------------- Droits
    def test_droits(self):
        for url in ["/api/rapports/", "/api/rapports/etat-stock/", "/api/historique/", "/api/historique/export/?fichier=pdf"]:
            self.assertEqual(self.a.get(url).status_code, 200, url)
            self.assertEqual(self.s.get(url).status_code, 403, url)
            self.assertEqual(self.r.get(url).status_code, 403, url)
            self.assertEqual(APIClient().get(url).status_code, 401, url)
        self.assertEqual(self.a.get("/api/rapports/inconnu/").status_code, 404)
        self.assertEqual(self.a.get("/api/rapports/etat-stock/?fichier=doc").status_code, 400)

    # ---------------------------------------------------------------- Rapports
    def test_catalogue(self):
        d = self.a.get("/api/rapports/").data
        self.assertEqual([r["cle"] for r in d["rapports"]], RAPPORTS)
        self.assertTrue(any(m["nom"] == "Ankorondrano" for m in d["magasins"]))
        self.assertNotIn("fonction", d["rapports"][0])

    def test_tous_les_rapports_ecran_pdf_excel(self):
        self.circuit()
        from openpyxl import load_workbook
        for cle in RAPPORTS:
            rep = self.a.get(f"/api/rapports/{cle}/")
            self.assertEqual(rep.status_code, 200, cle)
            self.assertTrue(rep.data["titre"] and rep.data["colonnes"], cle)
            for s in rep.data["sections"]:
                for l in s["lignes"]:
                    self.assertFalse([k for k in l if k in ("_jours", "_le")], cle)
            pdf = self.a.get(f"/api/rapports/{cle}/?fichier=pdf")
            self.assertEqual(pdf.status_code, 200, cle)
            self.assertTrue(pdf.content.startswith(b"%PDF"), cle)
            self.assertIn(".pdf", pdf["Content-Disposition"])
            xlsx = self.a.get(f"/api/rapports/{cle}/?fichier=xlsx")
            self.assertEqual(xlsx.status_code, 200, cle)
            classeur = load_workbook(BytesIO(xlsx.content))
            self.assertIn(rep.data["titre"], classeur.active["A1"].value)

    def test_etat_du_stock(self):
        d = self.a.get(f"/api/rapports/etat-stock/?magasin={self.ank.id}").data
        lignes = [l for s in d["sections"] for l in s["lignes"]]
        self.assertEqual(len(lignes), articles_avec_quantites().filter(emplacement__magasin=self.ank).count())
        c = next(l for l in lignes if l["code"] == self.conso.code)
        self.assertEqual((c["qte"], c["minimum"], c["etat"]), (50, 10, "OK"))

    def test_a_commander_et_valeur_comme_le_tableau_de_bord(self):
        Article.objects.filter(id=self.conso.id).update(qte_stock=3, qte_minimale=20)
        d = self.a.get("/api/rapports/a-commander/").data
        c = next(l for s in d["sections"] for l in s["lignes"] if l["code"] == self.conso.code)
        self.assertEqual(c["a_commander"], 17)
        self.assertEqual(Decimal(c["valeur"]), Decimal(17 * 2500))
        nb = self.a.get("/api/tableau-de-bord/").data["cartes"]["sous_seuil"]
        self.assertEqual(sum(len(s["lignes"]) for s in d["sections"]), nb)
        v = self.a.get("/api/rapports/valeur-stock/").data
        self.assertEqual(Decimal(v["total"]["valeur"]), valeur_du_stock().quantize(Decimal("1")))

    def test_consommation_et_materiel_prete(self):
        d = self.circuit(retour_prevu=timezone.localdate() - datetime.timedelta(days=0))
        conso = self.a.get("/api/rapports/consommation/").data
        self.assertEqual([s["titre"] for s in conso["sections"]], ["Fibre Ivato"])
        ligne = conso["sections"][0]["lignes"][0]
        self.assertEqual((ligne["code"], ligne["qte"]), (self.conso.code, 4))
        self.assertIn(d.requisition.numero, ligne["bons"])
        self.assertEqual(Decimal(conso["total"]["valeur"]), Decimal(4 * 2500))
        par_service = self.a.get("/api/rapports/consommation/?regrouper=service").data
        self.assertEqual(par_service["titre"], "Consommation par service demandeur")
        avant = self.a.get("/api/rapports/consommation/?periode=mois_dernier").data
        self.assertEqual(avant["sections"], [])
        prete = self.a.get("/api/rapports/materiel-prete/").data
        lignes = [l for l in prete["sections"][0]["lignes"] if l["bon"].startswith("RM-")]  # (les exemples ont aussi du matériel « Prêté » sans bon)
        self.assertEqual(sorted(l["qte"] for l in lignes), [1, 1, 3])  # 3 casques + 2 n° de série
        self.assertTrue(all(l["prete_a"] == "Jean Rabe" and l["chantier"] == "Fibre Ivato" for l in lignes))

    def test_materiel_en_retard(self):
        d = self.circuit()
        Demande.objects.get(id=d.id).lignes.filter(article=self.casque).update(
            date_retour_prevue=timezone.localdate() - datetime.timedelta(days=3))
        lignes = self.a.get("/api/rapports/materiel-prete/").data["sections"][0]["lignes"]
        self.assertEqual(lignes[0]["code"], self.casque.code)  # les retards d'abord
        self.assertEqual((lignes[0]["retard"], lignes[0]["_ton_retard"]), ("3 j", "rouge"))

    def test_etalonnages(self):
        auj = timezone.localdate()
        Exemplaire.objects.filter(numero_serie="RP-SN-0").update(date_etalonnage=auj - datetime.timedelta(days=2))
        Exemplaire.objects.filter(numero_serie="RP-SN-1").update(date_etalonnage=auj + datetime.timedelta(days=200))
        lignes = self.a.get("/api/rapports/etalonnages/").data["sections"][0]["lignes"]
        series = [l["serie"] for l in lignes]
        self.assertIn("RP-SN-0", series)
        self.assertNotIn("RP-SN-1", series)  # au-delà de 3 mois
        self.assertEqual(next(l for l in lignes if l["serie"] == "RP-SN-0")["echeance"], "Dépassé de 2 j")
        tout = self.a.get("/api/rapports/etalonnages/?horizon=tout").data["sections"][0]["lignes"]
        self.assertIn("RP-SN-1", [l["serie"] for l in tout])

    # ---------------------------------------------------------------- Historique
    def test_historique_stock_avant_apres(self):
        d = self.circuit()
        rep = self.a.get("/api/historique/")
        self.assertEqual(rep.status_code, 200)
        lignes = rep.data["lignes"]
        types = [l["type"] for l in lignes]
        for t in ("entree", "sortie", "demande", "validation", "visa", "impression"):
            self.assertIn(t, types)
        # Les plus récentes d'abord
        self.assertEqual([l["le"] for l in lignes], sorted([l["le"] for l in lignes], reverse=True))
        # Consommable : 50 + 6 reçus = 56, puis 4 sortis = 52 (stock actuel)
        entree = next(l for l in lignes if l["type"] == "entree" and l["code"] == self.conso.code)
        sortie = next(l for l in lignes if l["type"] == "sortie" and l["code"] == self.conso.code)
        self.assertEqual((entree["quantite"], entree["stock"]), ("+6", "50 → 56"))
        self.assertEqual((sortie["quantite"], sortie["stock"]), ("−4", "56 → 52"))
        self.assertEqual(Article.objects.get(id=self.conso.id).qte_stock, 52)
        self.assertEqual(sortie["reference"], d.requisition.numero)
        self.assertEqual(sortie["par"], "Hery Rasoa")
        # Les 2 n° de série sortis ensemble : une seule ligne, −2
        serie = [l for l in lignes if l["type"] == "sortie" and l["code"] == self.serie.code]
        self.assertEqual(len(serie), 1)
        self.assertEqual(serie[0]["quantite"], "−2")
        self.assertIn("RP-SN-0", serie[0]["detail"])

    def test_historique_filtres_et_exports(self):
        self.circuit()
        self.assertEqual({l["type"] for l in self.a.get("/api/historique/?type=sortie").data["lignes"]}, {"sortie"})
        par = self.a.get(f"/api/historique/?utilisateur={self.saisisseur.id}").data["lignes"]
        self.assertEqual({l["type"] for l in par}, {"demande"})
        cherche = self.a.get("/api/historique/?q=casque chantier").data["lignes"]
        self.assertTrue(cherche and all("Casque" in l["article"] for l in cherche))
        self.assertEqual(self.a.get("/api/historique/?periode=mois_dernier").data["count"], 0)
        hier = (timezone.localdate() - datetime.timedelta(days=1)).isoformat()
        self.assertEqual(self.a.get(f"/api/historique/?periode=choisie&du={hier}&au={hier}").data["count"], 0)
        self.assertGreater(self.a.get(f"/api/historique/?periode=choisie&du={hier}").data["count"], 0)
        self.assertTrue(any(u["nom"] == "Jean Rabe" for u in self.a.get("/api/historique/").data["utilisateurs"]))
        pdf = self.a.get("/api/historique/export/?fichier=pdf&type=sortie")
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(self.a.get("/api/historique/export/?fichier=xlsx").content)).active
        valeurs = [c.value for ligne in ws.iter_rows() for c in ligne if c.value]
        self.assertIn("Stock avant → après", valeurs)
        self.assertIn("56 → 52", valeurs)

    def test_correction_dans_l_administration(self):
        """Une quantité corrigée dans l'administration est notée : le « stock avant → après » reste juste."""
        self.circuit()
        c = APIClient(); c.force_login(self.admin)
        a = Article.objects.get(id=self.conso.id)
        donnees = {"nom": a.nom, "categorie": a.categorie_id, "type": a.type, "unite": a.unite, "qte_stock": 40,
                   "qte_minimale": a.qte_minimale, "emplacement": a.emplacement_id, "prix_unitaire": "2500",
                   "stocks_etat-TOTAL_FORMS": 0, "stocks_etat-INITIAL_FORMS": 0,
                   "exemplaires-TOTAL_FORMS": 0, "exemplaires-INITIAL_FORMS": 0}
        rep = c.post(f"/admin/stock/article/{a.id}/change/", donnees)
        self.assertEqual(rep.status_code, 302, rep.content[:3000])
        m = MouvementEtat.objects.filter(type="CORRECTION", article=a).latest("id")
        self.assertEqual((m.variation, m.par), (-12, self.admin))
        lignes = self.a.get(f"/api/historique/?q={self.conso.code}").data["lignes"]
        self.assertEqual(lignes[0]["stock"], "52 → 40")
        self.assertEqual(next(l for l in lignes if l["type"] == "entree")["stock"], "50 → 56")
